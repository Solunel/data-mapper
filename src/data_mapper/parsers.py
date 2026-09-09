"""CSV/XLSX 的不可变本地接入与表格解析。"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import uuid
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook

from .contracts import (
    ParsedRecord,
    ParsedTable,
    PipelineConfig,
    RawDataset,
    SourceContextCell,
    resolved_path,
)
from .errors import InputParseError, UnsupportedFormatError
from .table_structure import (
    MergedRange,
    expand_header_rows,
    find_header_rows,
    flatten_headers,
    is_empty,
    row_is_empty,
    table_column_bounds,
)


_SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}


def ingest_raw(path: str | Path, config: PipelineConfig) -> RawDataset:
    source = resolved_path(path)
    if not source.is_file():
        raise InputParseError(f"输入不是普通文件：{source}")
    extension = source.suffix.lower()
    if extension not in _SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(
            f"不支持的输入格式“{extension or '<无扩展名>'}”；Phase 1 仅支持 .xlsx 和 .csv"
        )

    payload = source.read_bytes()
    checksum = hashlib.sha256(payload).hexdigest()
    dataset_id = str(uuid.uuid5(uuid.NAMESPACE_URL, source.as_uri()))
    config_json = json.dumps(config.to_dict(), sort_keys=True, ensure_ascii=False)
    version_id = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{dataset_id}:{checksum}:{config_json}"))
    return RawDataset(
        dataset_id=dataset_id,
        version_id=version_id,
        source_file=source.name,
        extension=extension,
        size_bytes=len(payload),
        sha256=checksum,
        content_ref=str(source),
        ingested_at=datetime.now(timezone.utc),
        parser_config=config.to_dict(),
        parser_version=config.parser_version,
        status="ready",
    )


def parse_raw(raw: RawDataset, config: PipelineConfig) -> tuple[ParsedTable, ...]:
    payload = Path(raw.content_ref).read_bytes()
    actual_checksum = hashlib.sha256(payload).hexdigest()
    if actual_checksum != raw.sha256:
        raise InputParseError(
            "Raw 文件在接入后发生变化，内容校验值不再匹配",
            raw_dataset=replace(raw, status="failed", error="checksum_mismatch"),
        )

    try:
        if raw.extension == ".csv":
            return (_parse_csv(payload, raw.source_file, config),)
        if raw.extension == ".xlsx":
            return _parse_xlsx(payload, raw.source_file, config)
    except InputParseError:
        raise
    except Exception as exc:
        message = f"无法解析 {raw.source_file}：{exc}"
        raise InputParseError(
            message,
            raw_dataset=replace(raw, status="failed", error=message),
        ) from exc
    raise UnsupportedFormatError(f"不支持的输入格式“{raw.extension}”")


def _decode_csv(payload: bytes, encoding: str | None) -> tuple[str, str]:
    candidates = (encoding,) if encoding else ("utf-8-sig", "gb18030")
    failures: list[str] = []
    for candidate in candidates:
        if candidate is None:
            continue
        try:
            return payload.decode(candidate, errors="strict"), candidate
        except (LookupError, UnicodeDecodeError) as exc:
            failures.append(f"{candidate}: {exc}")
    raise InputParseError("CSV 无法使用允许的编码解码：" + "; ".join(failures))


def _parse_csv(payload: bytes, source_file: str, config: PipelineConfig) -> ParsedTable:
    text, encoding = _decode_csv(payload, config.csv_encoding)
    delimiter = config.csv_delimiter
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = [tuple(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    warning = ()
    if encoding == "gb18030" and config.csv_encoding is None:
        warning = (f"CSV 使用受限回退编码 {encoding} 解码",)
    return _table_from_rows(source_file, "CSV", rows, config, warning, ())


def _parse_xlsx(payload: bytes, source_file: str, config: PipelineConfig) -> tuple[ParsedTable, ...]:
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=False, data_only=True)
    except Exception as exc:
        raise InputParseError(f"XLSX 文件无效或不受支持：{exc}") from exc

    tables: list[ParsedTable] = []
    try:
        for worksheet in workbook.worksheets:
            rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
            if not any(not row_is_empty(row) for row in rows):
                continue
            merged_ranges = tuple(
                (item.min_row, item.max_row, item.min_col, item.max_col)
                for item in worksheet.merged_cells.ranges
            )
            tables.append(
                _table_from_rows(
                    source_file,
                    worksheet.title,
                    rows,
                    config,
                    (),
                    merged_ranges,
                )
            )
    finally:
        workbook.close()
    if not tables:
        raise InputParseError("XLSX 不包含任何非空工作表")
    return tuple(tables)


def _table_from_rows(
    source_file: str,
    sheet_name: str,
    rows: Sequence[Sequence[Any]],
    config: PipelineConfig,
    warning_messages: Iterable[str],
    merged_ranges: Sequence[MergedRange],
) -> ParsedTable:
    last_column = max(
        (index + 1 for row in rows for index, value in enumerate(row) if not is_empty(value)),
        default=0,
    )
    if last_column == 0:
        raise InputParseError(f"{source_file} [{sheet_name}] 为空")
    bounded_rows = [tuple(row[:last_column]) for row in rows]
    header_rows = find_header_rows(
        bounded_rows,
        config.header_rows.get(sheet_name),
        merged_ranges,
    )
    header_row = header_rows[0]
    header_end_row = header_rows[-1]
    expanded = expand_header_rows(
        bounded_rows,
        header_rows,
        last_column,
        merged_ranges,
    )
    all_headers = flatten_headers(expanded, last_column)
    first_column, last_table_column = table_column_bounds(
        all_headers,
        bounded_rows[header_end_row:],
    )
    headers = tuple(all_headers[first_column:last_table_column])
    if not any(headers):
        raise InputParseError(f"{source_file} [{sheet_name}] 没有可用表头")

    context_cells = tuple(
        SourceContextCell(source_row=row_index, source_column=column_index, value=value)
        for row_index, row in enumerate(bounded_rows[: header_row - 1], start=1)
        for column_index, value in enumerate(row, start=1)
        if not is_empty(value)
    )
    records: list[ParsedRecord] = []
    for source_row, row in enumerate(
        bounded_rows[header_end_row:],
        start=header_end_row + 1,
    ):
        values = list(row)
        values.extend([None] * (last_column - len(values)))
        records.append(
            ParsedRecord(
                source_row=source_row,
                values=tuple(values[first_column:last_table_column]),
            )
        )

    from .contracts import Issue

    messages = list(warning_messages)
    if len(header_rows) > 1:
        messages.append(f"使用第 {header_rows[0]}-{header_rows[-1]} 行组合多行表头")
    ignored_columns = last_column - (last_table_column - first_column)
    if ignored_columns:
        messages.append(f"已忽略表格区域外的 {ignored_columns} 个边缘列")
    warnings = tuple(
        Issue(
            severity="warning",
            code="parse_warning",
            message=message,
            step="Parse",
            source_file=source_file,
            sheet_name=sheet_name,
        )
        for message in messages
    )
    return ParsedTable(
        source_file=source_file,
        sheet_name=sheet_name,
        header_row=header_row,
        header_rows=header_rows,
        original_headers=headers,
        source_column_positions=tuple(range(first_column + 1, last_table_column + 1)),
        context_cells=context_cells,
        records=tuple(records),
        warnings=warnings,
    )
