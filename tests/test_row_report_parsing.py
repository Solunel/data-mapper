from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from data_mapper import PipelineConfig, curate_file


def _new_workbook(path: Path, sheet_name: str = "指标在行"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    return workbook, sheet


def test_single_header_preserves_business_names_and_new_columns(tmp_path: Path) -> None:
    source = tmp_path / "synthetic_row_report.xlsx"
    workbook, sheet = _new_workbook(source)
    sheet.merge_cells("D1:I1")
    sheet["D1"] = "SYNTHETIC 财务快报"
    sheet.append(
        [
            "时间",
            "公司名称",
            "上级公司",
            "项      目",
            "行次",
            "本月数",
            "本年累计数",
            "新增预测数",
            "新增空列",
        ]
    )
    sheet.append(["2026年01月", "模拟公司", None, "销售商品、提供劳务收到的现金", 1, 12345678, None, 80.5, None])
    sheet.append(["2026年01月", "模拟公司", None, "经营活动现金流入小计", 2, "－", None, None, None])
    workbook.save(source)
    workbook.close()

    curated = curate_file(source).curated_datasets[0]

    expected_columns = [
        "时间",
        "公司名称",
        "上级公司",
        "项      目",
        "行次",
        "本月数",
        "本年累计数",
        "新增预测数",
        "新增空列",
    ]
    assert curated.header_row == 2
    assert curated.header_rows == (2,)
    assert [item.normalized_name for item in curated.header_mapping] == expected_columns
    assert [item.source_position for item in curated.header_mapping] == list(range(1, 10))
    assert [row.source_row for row in curated.rows] == [3, 4]
    assert list(curated.rows[0].values) == expected_columns
    assert curated.rows[0].values["项      目"] == "销售商品、提供劳务收到的现金"
    assert curated.rows[0].values["本月数"] == 12345678
    assert curated.rows[1].values["本月数"] is None
    assert curated.rows[0].values["新增预测数"] == 80.5
    assert curated.data_schema.columns[-2].data_type == "number"
    assert curated.data_schema.columns[-1].data_type == "null"


def test_two_level_header_and_left_margin_are_parsed_without_business_reshaping(
    tmp_path: Path,
) -> None:
    source = tmp_path / "synthetic_two_level_report.xlsx"
    workbook, sheet = _new_workbook(source, "资产负债结构")
    sheet.merge_cells("B2:F2")
    sheet["B2"] = "SYNTHETIC 资产负债表"
    sheet["E4"] = "2026年01月"
    sheet["F4"] = "单位：元"
    sheet.merge_cells("B5:B6")
    sheet["B5"] = "科目名称"
    sheet.merge_cells("C5:C6")
    sheet["C5"] = "行次"
    sheet.merge_cells("D5:E5")
    sheet["D5"] = "增减额"
    sheet["D6"] = "比年初"
    sheet["E6"] = "比上年同期"
    sheet.merge_cells("F5:F6")
    sheet["F5"] = "新增预算数"
    sheet.append([None, "货币资金", 1, 100, 90, None])
    sheet.append([None, "应收账款", 2, -20, "－", None])
    workbook.save(source)
    workbook.close()

    curated = curate_file(source).curated_datasets[0]

    expected_columns = [
        "科目名称",
        "行次",
        "增减额 / 比年初",
        "增减额 / 比上年同期",
        "新增预算数",
    ]
    assert curated.header_rows == (5, 6)
    assert [item.normalized_name for item in curated.header_mapping] == expected_columns
    assert [item.source_position for item in curated.header_mapping] == [2, 3, 4, 5, 6]
    assert [row.source_row for row in curated.rows] == [7, 8]
    assert list(curated.rows[0].values) == expected_columns
    assert curated.rows[0].values["科目名称"] == "货币资金"
    assert curated.rows[1].values["增减额 / 比上年同期"] is None
    assert {(cell.source_row, cell.source_column, cell.value) for cell in curated.context_cells} == {
        (2, 2, "SYNTHETIC 资产负债表"),
        (4, 5, "2026年01月"),
        (4, 6, "单位：元"),
    }


def test_three_level_merged_header_keeps_added_value_columns(tmp_path: Path) -> None:
    source = tmp_path / "synthetic_three_level_report.xlsx"
    workbook, sheet = _new_workbook(source, "成本结构")
    sheet.merge_cells("B2:G2")
    sheet["B2"] = "SYNTHETIC 成本费用表"
    sheet.merge_cells("B5:B7")
    sheet["B5"] = "费用明细"
    sheet.merge_cells("C5:C7")
    sheet["C5"] = "行次"
    sheet.merge_cells("D5:G5")
    sheet["D5"] = "本年金额"
    sheet.merge_cells("D6:E6")
    sheet["D6"] = "生产成本"
    sheet.merge_cells("F6:G6")
    sheet["F6"] = "期间费用"
    sheet["D7"] = "发电成本"
    sheet["E7"] = "购电成本"
    sheet["F7"] = "管理费用"
    sheet["G7"] = "新增销售费用"
    sheet.append([None, "折旧费", 1, 10, 20, 30, 40])
    workbook.save(source)
    workbook.close()

    curated = curate_file(source).curated_datasets[0]

    assert curated.header_rows == (5, 6, 7)
    assert [item.normalized_name for item in curated.header_mapping] == [
        "费用明细",
        "行次",
        "本年金额 / 生产成本 / 发电成本",
        "本年金额 / 生产成本 / 购电成本",
        "本年金额 / 期间费用 / 管理费用",
        "本年金额 / 期间费用 / 新增销售费用",
    ]
    assert len(curated.rows) == 1
    assert len(curated.rows[0].values) == 6
    assert curated.rows[0].values["本年金额 / 期间费用 / 新增销售费用"] == 40


def test_non_first_header_csv_accepts_added_columns_without_fixed_business_names(
    tmp_path: Path,
) -> None:
    source = tmp_path / "synthetic_row_report.csv"
    source.write_text(
        "SYNTHETIC 月度明细,,,,\n"
        "费用明细,序号,实际发生额,预算金额,差异率\n"
        "人工成本,1,12345678,12000000,2.88%\n",
        encoding="utf-8",
    )

    curated = curate_file(source).curated_datasets[0]

    assert curated.header_rows == (2,)
    assert [item.normalized_name for item in curated.header_mapping] == [
        "费用明细",
        "序号",
        "实际发生额",
        "预算金额",
        "差异率",
    ]
    assert curated.rows[0].values["实际发生额"] == 12345678
    assert curated.rows[0].values["差异率"] == "2.88%"


def test_unreliable_explicit_header_is_visible_in_quality_report(tmp_path: Path) -> None:
    source = tmp_path / "synthetic_wrong_header.csv"
    source.write_text(
        "SYNTHETIC 指标报表,,,,,,\n"
        ",,,2026年01月,,,单位：元\n"
        "费用明细,序号,本期数,预算数,同期数,增减额,增减率\n"
        "人工成本,1,100,90,80,20,25%\n",
        encoding="utf-8",
    )

    curated = curate_file(
        source,
        PipelineConfig(header_rows={"CSV": 2}),
    ).curated_datasets[0]

    assert not curated.quality_report.passed
    assert any(issue.code == "excessive_empty_headers" for issue in curated.issues)
